"""Parser for Household Financial Savings (Bulletin Table 50A).

Source: RBI Bulletin — "Flow of Financial Assets and Liabilities of Households"
3 sheets: FY 2022-23, FY 2023-24, FY 2024-25
Each sheet has quarterly (Q1-Q4) + annual data.
"""

import json
import os

import openpyxl


def _sanitize(obj):
    if isinstance(obj, dict):
        return {k: _sanitize(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_sanitize(v) for v in obj]
    if isinstance(obj, float) and (obj != obj or obj == float("inf")):
        return None
    return obj


ROW_MAP = {
    "net financial assets (i-ii)": "net_financial_assets",
    "per cent of gdp": None,  # skip — we'll extract it separately
    "i. financial assets": "total_financial_assets",
    "1.total deposits (a+b)": "deposits_total",
    "(a) bank deposits": "deposits_bank",
    "(b) non-bank deposits": "deposits_nonbank",
    "2. life insurance funds": "life_insurance",
    "3. provident and pension funds (including ppf)": "provident_pension",
    "4. currency": "currency",
    "5. investments": "investments_total",
    "(a) mutual funds": "mutual_funds",
    "(b) equity": "equity",
    "6. small savings (excluding ppf)": "small_savings",
    "ii. financial liabilities": "total_liabilities",
}


def parse(filepath: str) -> list[dict]:
    """Parse the Household Savings XLSX and return a list of records."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    records = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Extract FY from row 4
        fy_cell = ws.cell(row=4, column=3).value
        if not fy_cell:
            continue
        fy = str(fy_cell).strip()  # e.g. "2022-23"

        # Column mapping: C=Q1, D=Q2, E=Q3, F=Q4, G=Annual
        quarters = {
            3: f"{fy} Q1",
            4: f"{fy} Q2",
            5: f"{fy} Q3",
            6: f"{fy} Q4",
            7: f"{fy} Annual",
        }

        # Also track GDP % rows
        gdp_pct_for = None  # which item the next "Per cent of GDP" belongs to

        for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=False):
            label_raw = row[1].value
            if not label_raw:
                continue
            label = str(label_raw).strip().lower()

            # Check if this is a "Per cent of GDP" row
            if label == "per cent of gdp" and gdp_pct_for:
                for col_idx, period in quarters.items():
                    val = row[col_idx - 1].value
                    if val is not None:
                        try:
                            records.append({
                                "period": period,
                                "item": f"{gdp_pct_for}_pct_gdp",
                                "value": float(val),
                            })
                        except (ValueError, TypeError):
                            pass
                gdp_pct_for = None
                continue

            # Match to our row map
            key = ROW_MAP.get(label)
            if key is None:
                # Check partial matches
                for map_label, map_key in ROW_MAP.items():
                    if map_label in label and map_key:
                        key = map_key
                        break

            if not key:
                continue

            gdp_pct_for = key  # next "Per cent of GDP" row belongs to this item

            for col_idx, period in quarters.items():
                val = row[col_idx - 1].value
                if val is not None:
                    try:
                        records.append({
                            "period": period,
                            "item": key,
                            "value": float(val),
                        })
                    except (ValueError, TypeError):
                        pass

    return records


def export_json(records: list[dict], output_dir: str = "web/public/data"):
    os.makedirs(output_dir, exist_ok=True)

    with open(os.path.join(output_dir, "household_savings.json"), "w") as f:
        json.dump(_sanitize(records), f, indent=2)
    print(f"  Exported household_savings.json ({len(records)} records)")

    # Also create a summary for the chart — annual data by instrument
    annual = [r for r in records if "Annual" in r["period"]]
    instruments = [
        "deposits_bank", "life_insurance", "provident_pension",
        "currency", "mutual_funds", "equity", "small_savings",
    ]
    summary = []
    for r in annual:
        if r["item"] in instruments:
            fy = r["period"].replace(" Annual", "")
            summary.append({"fy": fy, "instrument": r["item"], "value_crore": r["value"]})

    with open(os.path.join(output_dir, "household_savings_summary.json"), "w") as f:
        json.dump(_sanitize(summary), f, indent=2)
    print(f"  Exported household_savings_summary.json ({len(summary)} records)")

    # GDP % summary
    gdp_pct = [r for r in records if "_pct_gdp" in r["item"] and "Annual" in r["period"]]
    with open(os.path.join(output_dir, "household_savings_gdp.json"), "w") as f:
        json.dump(_sanitize(gdp_pct), f, indent=2)
    print(f"  Exported household_savings_gdp.json ({len(gdp_pct)} records)")


def parse_html(filepath: str) -> list[dict]:
    """Parse a Bulletin HTML page containing Table 50(a)/52(a)."""
    import pandas as pd
    from io import BytesIO

    with open(filepath, "rb") as f:
        tables = pd.read_html(BytesIO(f.read()))

    records = []
    for table in tables:
        if table.shape[0] < 10 or table.shape[1] < 6:
            continue

        # Find the FY from header rows
        fy = None
        for r in range(min(5, len(table))):
            for c in range(table.shape[1]):
                val = str(table.iloc[r, c]).strip()
                if re.match(r"\d{4}-\d{2}", val):
                    fy = val
                    break
            if fy:
                break

        if not fy:
            continue

        quarters = {1: f"{fy} Q1", 2: f"{fy} Q2", 3: f"{fy} Q3", 4: f"{fy} Q4", 5: f"{fy} Annual"}

        gdp_pct_for = None
        for r in range(len(table)):
            label_raw = table.iloc[r, 0]
            if not label_raw or str(label_raw).strip() == "" or str(label_raw) == "nan":
                continue
            label = str(label_raw).strip().lower()

            if label == "per cent of gdp" and gdp_pct_for:
                for col_idx, period in quarters.items():
                    val = table.iloc[r, col_idx]
                    try:
                        records.append({"period": period, "item": f"{gdp_pct_for}_pct_gdp", "value": float(val)})
                    except (ValueError, TypeError):
                        pass
                gdp_pct_for = None
                continue

            key = ROW_MAP.get(label)
            if key is None:
                for map_label, map_key in ROW_MAP.items():
                    if map_label in label and map_key:
                        key = map_key
                        break
            if not key:
                continue

            gdp_pct_for = key
            for col_idx, period in quarters.items():
                val = table.iloc[r, col_idx]
                try:
                    records.append({"period": period, "item": key, "value": float(val)})
                except (ValueError, TypeError):
                    pass

    return records


if __name__ == "__main__":
    import re

    # Parse current XLSX (FY 2022-23 to 2024-25)
    records = parse(
        "web/public/data/raw/household_savings/50AT_BUL28082025C459E5FE597241EA9BBB5BD5FB5B0BAF.XLSX"
    )

    # Parse older HTML (FY 2019-20 to 2021-22)
    older = parse_html(
        "web/public/data/raw/household_savings/bulletin_t50a_oct2022_id21387.html"
    )
    records = older + records

    # Deduplicate (prefer newer data)
    seen = set()
    deduped = []
    for r in reversed(records):
        key = (r["period"], r["item"])
        if key not in seen:
            seen.add(key)
            deduped.append(r)
    deduped.reverse()

    print(f"Total: {len(deduped)} records")
    for period in sorted(set(r["period"] for r in deduped)):
        items = [r for r in deduped if r["period"] == period]
        print(f"  {period}: {len(items)} items")

    export_json(deduped)

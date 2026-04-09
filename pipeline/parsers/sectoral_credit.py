"""Parser for Sectoral Deployment of Bank Credit.

Source: RBI press release XLSX — Statement 1
Has 5 date snapshots with outstanding amounts and YoY growth.
Sectors: Agriculture, Industry (Micro/Small/Medium/Large), Services, Personal Loans.
"""

import json
import os

import openpyxl


def _sanitize(obj):
    if isinstance(obj, dict):
        return {k: _sanitize(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_sanitize(v) for v in obj]
    if isinstance(obj, float) and (obj != obj or obj == float("inf")):
        return None
    return obj


def _to_float(v):
    if v is None:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


SECTOR_MAP = {
    "i. bank credit (ii + iii)": "total_bank_credit",
    "ii. food credit": "food_credit",
    "iii. non-food credit": "non_food_credit",
    "1. agriculture and allied activities": "agriculture",
    "2. industry (micro and small, medium and large)": "industry_total",
    "2.1. micro and small": "industry_micro_small",
    "2.2. medium": "industry_medium",
    "2.3. large": "industry_large",
    "3. services": "services",
    "4. personal loans": "personal_loans",
    "4.1. consumer durables": "personal_consumer_durables",
    "4.2. housing": "personal_housing",
    "4.3. advances against fixed deposits": "personal_fd",
    "4.5. credit card outstanding": "personal_credit_card",
    "4.6. education": "personal_education",
    "4.7. vehicle loans": "personal_vehicle",
    "4.8. loans against gold jewellery": "personal_gold",
    "4.9. other personal loans": "personal_other",
}


def parse(filepath: str) -> list[dict]:
    """Parse the Sectoral Credit XLSX Statement 1."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb["Statement 1"]

    # Extract date columns from row 4
    dates = []
    for col in range(2, ws.max_column + 1):
        val = ws.cell(row=4, column=col).value
        if val and "." in str(val) and ("20" in str(val) or "Mar" in str(val)):
            dates.append((col, str(val).strip()))

    # Only take the first 5 (outstanding amounts, not % change columns)
    date_cols = dates[:5]

    records = []
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=False):
        label_raw = row[0].value
        if not label_raw:
            continue
        label = str(label_raw).strip().lower()

        key = None
        for map_label, map_key in SECTOR_MAP.items():
            if label.startswith(map_label) or map_label in label:
                key = map_key
                break

        if not key:
            continue

        for col_idx, date_label in date_cols:
            val = _to_float(row[col_idx - 1].value)
            if val is not None:
                records.append({
                    "date": date_label,
                    "sector": key,
                    "outstanding_crore": round(val, 2),
                })

    return records


def export_json(records: list[dict], output_dir: str = "web/public/data"):
    os.makedirs(output_dir, exist_ok=True)

    with open(os.path.join(output_dir, "sectoral_credit.json"), "w") as f:
        json.dump(_sanitize(records), f, indent=2)
    print(f"  Exported sectoral_credit.json ({len(records)} records)")

    # Summary: latest date, main 4 sectors
    dates = sorted(set(r["date"] for r in records))
    latest = dates[-1]
    main_sectors = ["agriculture", "industry_total", "services", "personal_loans"]
    summary = [r for r in records if r["date"] == latest and r["sector"] in main_sectors]
    with open(os.path.join(output_dir, "sectoral_credit_latest.json"), "w") as f:
        json.dump(_sanitize(summary), f, indent=2)
    print(f"  Exported sectoral_credit_latest.json ({len(summary)} records, date={latest})")


if __name__ == "__main__":
    records = parse("web/public/data/raw/sectoral_credit/SIBC30032026.xlsx")
    print(f"Total: {len(records)} records")
    dates = sorted(set(r["date"] for r in records))
    print(f"Dates: {dates}")
    sectors = sorted(set(r["sector"] for r in records))
    print(f"Sectors: {sectors}")

    # Show latest snapshot
    latest = dates[-1]
    print(f"\nLatest ({latest}):")
    for r in sorted([r for r in records if r["date"] == latest], key=lambda x: -x["outstanding_crore"]):
        print(f"  {r['sector']:30s}: ₹{r['outstanding_crore']/100000:>8.1f} lakh crore")

    export_json(records)
